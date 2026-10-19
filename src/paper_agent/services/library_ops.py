"""文献库管理操作：本地文件导入与论文删除（CLI 与 Web 共用）。

删除是「全链清理」：账本记录、向量库块、解析/知识/原文产物一并移除——
留一半状态会让下次检索/问答出现幽灵数据。
"""

from __future__ import annotations

import hashlib
import re
import shutil
import uuid
from pathlib import Path

from paper_agent import paths
from paper_agent.download import safe_dirname
from paper_agent.library import Library, PaperRecord
from paper_agent.parsing import SUPPORTED_SUFFIXES
from paper_agent.rag.vectorstore import VectorStore
from paper_agent.services.locks import paper_lock


class UnsupportedFormat(ValueError):
    pass


class CleanupPending(RuntimeError):
    pass


def validate_document(path: Path) -> None:
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormat("仅支持 PDF/DOCX 文件")
    try:
        if path.suffix.lower() == ".pdf":
            import pymupdf
            if not path.read_bytes()[:1024].lstrip(b"\xef\xbb\xbf \t\r\n").startswith(b"%PDF-"):
                raise ValueError("文件内容不是 PDF")
            with pymupdf.open(path) as doc:
                if doc.needs_pass or not doc.page_count:
                    raise ValueError("PDF 需要密码或没有页面，请换用可读取的文件")
                if not any(page.get_text().strip() or page.get_images() or page.get_drawings() or page.first_annot for page in doc):
                    raise ValueError("PDF 所有页面均为空白，请更换含正文的文件")
        else:
            from docx import Document
            Document(str(path))
    except Exception as exc:
        raise UnsupportedFormat(f"文件无法解析，可能已损坏或格式不符；请更换文件：{exc}") from exc


def attach_fulltext(path: Path, sid: str, *, cfg: dict, lib: Library, replace_existing: bool = False) -> PaperRecord:
    validate_document(path)
    with paper_lock(lib, sid):
        rec = lib.get(sid)
        if rec is None:
            raise ValueError("论文不存在")
        if rec.pdf_path and Path(rec.pdf_path).is_file() and not replace_existing:
            raise FileExistsError("此记录已有本地全文；为避免覆盖原文，未替换文件")
        root = paths.papers_dir(cfg).resolve()
        directory = (root / safe_dirname(sid)).resolve()
        if directory == root or not directory.is_relative_to(root):
            raise ValueError("全文目录超出文献目录")
        if replace_existing:
            # Validate the new document before invalidating old derived evidence.
            VectorStore(paths.vector_dir(cfg)).remove_paper(sid)
            for base in (paths.parsed_dir(cfg), paths.knowledge_dir(cfg)):
                artifact_dir = (base / safe_dirname(sid)).resolve()
                if artifact_dir == base.resolve() or not artifact_dir.is_relative_to(base.resolve()):
                    raise ValueError("产物目录超出文献目录")
                if artifact_dir.exists():
                    shutil.rmtree(artifact_dir)
        directory.mkdir(parents=True, exist_ok=True)
        dest = directory / f"paper{path.suffix.lower()}"
        temporary = directory / f"{uuid.uuid4().hex}.part"
        try:
            shutil.copyfile(path, temporary)
            temporary.replace(dest)
        finally:
            temporary.unlink(missing_ok=True)
        lib.set_status(sid, "downloaded", pdf_path=str(dest), error=None)
        return lib.get(sid)


def import_local_file(path: Path, *, cfg: dict, lib: Library, tag: str = "本地上传") -> tuple[PaperRecord, bool]:
    """导入本地 PDF/DOCX：内容哈希幂等，状态直达 downloaded。

    返回 (账本记录, 是否为新导入)；重复导入幂等合并，状态不变。
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormat(f"不支持的格式（{suffix}），仅支持 PDF/DOCX")
    validate_document(path)
    data = path.read_bytes()
    digest = hashlib.md5(data).hexdigest()[:8]
    stem = re.sub(r"[\W_]+", "-", path.stem, flags=re.UNICODE).strip("-")[:40] or "document"
    sid = f"local:{stem}-{digest}"

    dest_dir = paths.papers_dir(cfg) / safe_dirname(sid)
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"paper{suffix}"
    if not dest.exists() or dest.read_bytes() != data:
        dest.write_bytes(data)

    title = _title_from_document(path, fallback=stem.replace("-", " "))
    from paper_agent.models import Paper

    is_new = lib.upsert_paper(Paper(source="local", source_id=sid, title=title), tag=tag)
    if lib.get(sid) is not None and lib.get(sid).status == "discovered":
        lib.set_status(sid, "downloaded", error=None, pdf_path=str(dest))
    rec = lib.get(sid)
    assert rec is not None
    return rec, is_new


def delete_paper(source_id: str, *, cfg: dict, lib: Library) -> bool:
    """全链删除一篇论文：账本 + 向量块 + 原文/解析/知识产物目录。

    返回是否存在（不存在视为删除失败，由调用方决定 404）。
    """
    with paper_lock(lib, source_id):
        existed = lib.begin_delete(source_id)
        pending = source_id in {r["source_id"] for r in lib.pending_deletions()}
        if not existed and not pending:
            return False
        errors = []
        try:
            VectorStore(paths.vector_dir(cfg)).remove_paper(source_id)
        except Exception as exc:
            errors.append(f"索引清理失败：{type(exc).__name__}: {exc}")
        for root in (paths.parsed_dir(cfg), paths.knowledge_dir(cfg), paths.papers_dir(cfg), paths.data_dir(cfg) / "repro"):
            directory = (root / safe_dirname(source_id)).resolve()
            try:
                if directory == root.resolve() or not directory.is_relative_to(root.resolve()):
                    raise ValueError("清理目录超出文献目录")
                if directory.exists():
                    shutil.rmtree(directory)
            except Exception as exc:
                errors.append(f"文件清理失败：{type(exc).__name__}: {exc}")
        lib.finish_delete(source_id, "; ".join(errors))
        if errors:
            raise CleanupPending("已移除文献记录，部分产物待清理；可在任务记录中重试。" + "; ".join(errors))
        return True


def _title_from_document(path: Path, *, fallback: str) -> str:
    """取文档内建标题元数据；取不到则退回清洗后的文件名。"""
    if path.suffix.lower() == ".pdf":
        try:
            import pymupdf

            with pymupdf.open(path) as doc:
                meta_title = (doc.metadata or {}).get("title") or ""
            if meta_title.strip() and "untitled" not in meta_title.lower():
                return meta_title.strip()[:200]
        except Exception:
            pass
    elif path.suffix.lower() == ".docx":
        try:
            from docx import Document

            doc = Document(str(path))
            for para in doc.paragraphs[:8]:
                style = (para.style.name or "").lower() if para.style is not None else ""
                text = para.text.strip()
                if text and style in ("title", "heading 1"):
                    return text[:200]
        except Exception:
            pass
    return fallback
